"""Workbook read model (`WorkbookSnapshot`) and the openpyxl-backed `WorkbookDriver`."""

from __future__ import annotations

import csv
import io
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from core.manifest import ACTION_TYPES, Manifest

CELL_RE = re.compile(r"^(?P<sheet>[^!]+)!(?P<start>\$?[A-Z]+\$?[0-9]+)(?::(?P<end>\$?[A-Z]+\$?[0-9]+))?$")
_REF_RE = re.compile(r"(?<![A-Za-z0-9_\"])(\$?)([A-Z]{1,3})(\$?)([0-9]+)(?![0-9A-Za-z_(])")


@dataclass(frozen=True)
class CellFormat:
    background: str | None = None


@dataclass(frozen=True)
class ManifestSummary:
    applied: int
    errors: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class CellChange:
    sheet: str
    cell: str
    old: Any
    new: Any


@dataclass(frozen=True)
class ParsedRange:
    sheet: str
    start_col: int
    start_row: int
    end_col: int
    end_row: int


# ─── A1 helpers ──────────────────────────────────────────────────────────────


def column_index(letters: str) -> int:
    value = 0
    for char in letters.upper():
        value = value * 26 + (ord(char) - ord("A") + 1)
    return value - 1


def column_name(index: int) -> str:
    return get_column_letter(index + 1)


def split_cell(cell: str) -> tuple[int, int]:
    cell = cell.replace("$", "")
    letters = "".join(char for char in cell if char.isalpha())
    digits = "".join(char for char in cell if char.isdigit())
    return column_index(letters), int(digits) - 1


def parse_range(range_name: str) -> ParsedRange:
    match = CELL_RE.match(str(range_name).strip())
    if not match:
        raise ValueError(f"Invalid Calc range: {range_name}")
    start_col, start_row = split_cell(match.group("start"))
    end_col, end_row = split_cell(match.group("end") or match.group("start"))
    return ParsedRange(
        sheet=match.group("sheet"),
        start_col=min(start_col, end_col),
        start_row=min(start_row, end_row),
        end_col=max(start_col, end_col),
        end_row=max(start_row, end_row),
    )


def shift_formula_rows(formula: str, offset: int) -> str:
    """Shift relative row references (not `$`-anchored) by `offset` rows."""
    if offset == 0:
        return formula

    def repl(match: re.Match[str]) -> str:
        col_abs, col, row_abs, row = match.groups()
        if row_abs:
            return match.group(0)
        return f"{col_abs}{col}{row_abs}{max(int(row) + offset, 1)}"

    return _REF_RE.sub(repl, formula)


def coerce_value(value: str) -> Any:
    if value == "":
        return ""
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


# Backwards-compatible private aliases used by older modules/tests.
_parse_range = parse_range
_column_name = column_name
_column_index = column_index
_split_cell = split_cell
_coerce_value = coerce_value


# ─── Snapshot ────────────────────────────────────────────────────────────────


@dataclass
class WorkbookSnapshot:
    sheets: dict[str, list[list[Any]]]
    source: str = ""

    @classmethod
    def from_xlsx(cls, path: str, sheet_name: str | None = None, all_sheets: bool = False) -> WorkbookSnapshot:
        wb = openpyxl.load_workbook(path, data_only=False)
        if all_sheets:
            names = wb.sheetnames
        else:
            names = [sheet_name if sheet_name in wb.sheetnames else wb.sheetnames[0]]
        sheets = {name: _rows_from_worksheet(wb[name]) for name in names}
        return cls(sheets=sheets, source=str(path))

    @classmethod
    def from_file(cls, path: str, sheet_name: str = "Sheet1") -> WorkbookSnapshot:
        if str(path).lower().endswith((".xlsx", ".xlsm", ".xls")):
            return cls.from_xlsx(path, sheet_name=sheet_name)
        return cls.from_csv(path, sheet_name=sheet_name)

    @classmethod
    def from_csv(cls, path: str, sheet_name: str = "Sheet1") -> WorkbookSnapshot:
        with open(path, newline="", encoding="utf-8-sig") as handle:
            snapshot = cls.from_csv_text(handle.read(), sheet_name=sheet_name)
        snapshot.source = str(path)
        return snapshot

    @classmethod
    def from_csv_text(cls, text: str, sheet_name: str = "Sheet1") -> WorkbookSnapshot:
        rows = [[coerce_value(value) for value in row] for row in csv.reader(io.StringIO(text))]
        return cls(sheets={sheet_name: rows}, source="stdin")

    def dims(self, sheet_name: str) -> tuple[int, int]:
        rows = self.sheets.get(sheet_name, [])
        return len(rows), max((len(row) for row in rows), default=0)

    def data_range(self, sheet_name: str) -> str:
        n_rows, n_cols = self.dims(sheet_name)
        return f"{sheet_name}!A1:{column_name(max(n_cols - 1, 0))}{max(n_rows, 1)}"

    def to_csv(self, sheet_name: str = "Sheet1") -> str:
        buffer = io.StringIO()
        csv.writer(buffer).writerows(self.sheets.get(sheet_name, []))
        return buffer.getvalue()

    def save_csv(self, path: str, sheet_name: str = "Sheet1") -> None:
        with open(path, "w", newline="", encoding="utf-8") as handle:
            csv.writer(handle).writerows(self.sheets.get(sheet_name, []))

    def save_xlsx(self, path: str, sheet_name: str = "Sheet1", formats: dict[str, CellFormat] | None = None) -> None:
        """Compatibility helper: write one sheet with optional background colours."""
        driver = WorkbookDriver.from_snapshot(WorkbookSnapshot(sheets={sheet_name: self.sheets.get(sheet_name, [])}))
        for ref, fmt in (formats or {}).items():
            if fmt.background:
                driver._highlight_cells(parse_range(ref if "!" in ref else f"{sheet_name}!{ref}"), fmt.background)
        driver.workbook.save(path)

    def range_to_csv(self, range_name: str) -> str:
        buffer = io.StringIO()
        csv.writer(buffer).writerows(self._slice(range_name))
        return buffer.getvalue()

    def range_to_json(self, range_name: str) -> str:
        parsed = parse_range(range_name)
        rows = self._slice(range_name)
        if not rows or parsed.sheet not in self.sheets:
            return "[]"
        sheet = self.sheets[parsed.sheet]
        header_row = sheet[0][parsed.start_col : parsed.end_col + 1] if sheet else []
        headers = [str(header) for header in header_row]
        data_rows = rows if parsed.start_row > 0 else rows[1:]
        return json.dumps([dict(zip(headers, row, strict=False)) for row in data_rows])

    def _slice(self, range_name: str) -> list[list[Any]]:
        parsed = parse_range(range_name)
        sheet = self.sheets.get(parsed.sheet)
        if sheet is None:
            raise KeyError(f"Unknown sheet: {parsed.sheet}")
        result: list[list[Any]] = []
        for row in sheet[parsed.start_row : parsed.end_row + 1]:
            padded = row + [""] * (parsed.end_col + 1 - len(row))
            result.append(padded[parsed.start_col : parsed.end_col + 1])
        return result


def _rows_from_worksheet(ws: Any) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if value is None else value for value in row])
    # Trim trailing fully-empty rows that openpyxl reports because of formatting.
    while rows and all(value == "" for value in rows[-1]):
        rows.pop()
    return rows


def diff_snapshots(before: WorkbookSnapshot, after: WorkbookSnapshot) -> list[CellChange]:
    changes: list[CellChange] = []
    sheet_order = list(before.sheets) + [name for name in after.sheets if name not in before.sheets]
    for sheet in sheet_order:
        old_rows = before.sheets.get(sheet, [])
        new_rows = after.sheets.get(sheet, [])
        n_rows = max(len(old_rows), len(new_rows))
        n_cols = max([len(r) for r in old_rows] + [len(r) for r in new_rows] + [0])
        for r in range(n_rows):
            for c in range(n_cols):
                old = _cell_at(old_rows, r, c)
                new = _cell_at(new_rows, r, c)
                if old != new:
                    changes.append(CellChange(sheet=sheet, cell=f"{column_name(c)}{r + 1}", old=old, new=new))
    return changes


def _cell_at(rows: list[list[Any]], r: int, c: int) -> Any:
    if r >= len(rows) or c >= len(rows[r]):
        return None
    value = rows[r][c]
    return None if value == "" else value


# ─── Driver ──────────────────────────────────────────────────────────────────


class WorkbookDriver:
    """Applies manifests to an openpyxl workbook; saves to XLSX (full fidelity) or CSV (+ sidecar)."""

    def __init__(self, workbook: Any, active_sheet: str | None = None, source: str = ""):
        if isinstance(workbook, WorkbookSnapshot):
            built = WorkbookDriver.from_snapshot(workbook)
            workbook, active_sheet, source = built.workbook, built.active_sheet, built.source or source
        self.workbook = workbook
        self.active_sheet = active_sheet or workbook.sheetnames[0]
        self.source = source
        self._formats: dict[str, CellFormat] = {}

    # ── constructors ──

    @classmethod
    def from_snapshot(cls, snapshot: WorkbookSnapshot) -> WorkbookDriver:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name, rows in snapshot.sheets.items():
            ws = wb.create_sheet(title=name)
            for row in rows:
                ws.append([None if value == "" else value for value in row])
        if not wb.sheetnames:
            wb.create_sheet("Sheet1")
        return cls(wb, active_sheet=next(iter(snapshot.sheets), "Sheet1"), source=snapshot.source)

    @classmethod
    def from_csv_text(cls, text: str, sheet_name: str = "Sheet1") -> WorkbookDriver:
        return cls.from_snapshot(WorkbookSnapshot.from_csv_text(text, sheet_name=sheet_name))

    @classmethod
    def from_path(cls, path: str | Path, sheet_name: str | None = None) -> WorkbookDriver:
        path = Path(path)
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            wb = openpyxl.load_workbook(path)
            active = sheet_name if sheet_name in wb.sheetnames else wb.sheetnames[0]
            return cls(wb, active_sheet=active, source=str(path))
        text = path.read_text(encoding="utf-8-sig")
        driver = cls.from_csv_text(text, sheet_name=sheet_name or "Sheet1")
        driver.source = str(path)
        return driver

    @classmethod
    def from_stdin(cls, sheet_name: str = "Sheet1") -> WorkbookDriver:
        return cls.from_csv_text(sys.stdin.read(), sheet_name=sheet_name)

    # ── read model ──

    @property
    def sheet_names(self) -> list[str]:
        return list(self.workbook.sheetnames)

    @property
    def formats(self) -> dict[str, CellFormat]:
        return dict(self._formats)

    def snapshot(self) -> WorkbookSnapshot:
        return WorkbookSnapshot(
            sheets={name: _rows_from_worksheet(self.workbook[name]) for name in self.workbook.sheetnames},
            source=self.source,
        )

    # ── execution ──

    def execute_manifest(self, manifest: Manifest | dict[str, Any]) -> ManifestSummary:
        actions = manifest.model_dump()["actions"] if isinstance(manifest, Manifest) else manifest.get("actions", [])
        applied = 0
        errors: list[str] = []
        for action in actions:
            action_type = action.get("type")
            if action_type not in ACTION_TYPES:
                errors.append(f"Unsupported action type: {action_type}")
                continue
            try:
                getattr(self, f"_do_{action_type}")(action)
                applied += 1
            except Exception as exc:  # noqa: BLE001 - one bad action must not abort the rest
                errors.append(f"{action_type}: {exc}")
        return ManifestSummary(applied=applied, errors=errors)

    def _ws(self, sheet: str) -> Any:
        if sheet not in self.workbook.sheetnames:
            raise KeyError(f"Unknown sheet: {sheet}")
        return self.workbook[sheet]

    def _cells(self, parsed: ParsedRange):
        ws = self._ws(parsed.sheet)
        for r in range(parsed.start_row, parsed.end_row + 1):
            for c in range(parsed.start_col, parsed.end_col + 1):
                yield r, c, ws.cell(row=r + 1, column=c + 1)

    def _highlight_cells(self, parsed: ParsedRange, color: str) -> None:
        rgb = color.lstrip("#").upper()
        for r, c, cell in self._cells(parsed):
            cell.fill = PatternFill(fill_type="solid", fgColor=rgb, bgColor=rgb)
            self._formats[f"{parsed.sheet}!{column_name(c)}{r + 1}"] = CellFormat(background=color.lower())

    def _do_highlight(self, action: dict[str, Any]) -> None:
        self._highlight_cells(parse_range(action["range"]), str(action.get("color", "#f97316")))

    def _do_write_value(self, action: dict[str, Any]) -> None:
        parsed = parse_range(action["cell"])
        self._ws(parsed.sheet).cell(row=parsed.start_row + 1, column=parsed.start_col + 1).value = action.get("value")

    def _do_write_range(self, action: dict[str, Any]) -> None:
        parsed = parse_range(action["range"])
        ws = self._ws(parsed.sheet)
        for r_offset, row in enumerate(action["values"]):
            for c_offset, value in enumerate(row):
                ws.cell(row=parsed.start_row + 1 + r_offset, column=parsed.start_col + 1 + c_offset).value = value

    def _do_formula(self, action: dict[str, Any]) -> None:
        parsed = parse_range(action["cell"])
        formula = str(action["formula"])
        if not formula.startswith("="):
            formula = "=" + formula
        self._ws(parsed.sheet).cell(row=parsed.start_row + 1, column=parsed.start_col + 1).value = formula

    def _do_fill_formula(self, action: dict[str, Any]) -> None:
        parsed = parse_range(action["range"])
        template = str(action["formula"])
        if not template.startswith("="):
            template = "=" + template
        for r, c, cell in self._cells(parsed):
            if "{row}" in template or "{col}" in template:
                cell.value = template.replace("{row}", str(r + 1)).replace("{col}", column_name(c))
            else:
                cell.value = shift_formula_rows(template, r - parsed.start_row)

    def _do_set_format(self, action: dict[str, Any]) -> None:
        parsed = parse_range(action["range"])
        for r, c, cell in self._cells(parsed):
            font = cell.font
            kwargs: dict[str, Any] = {
                "name": font.name,
                "size": font.size,
                "bold": font.bold if action.get("bold") is None else action["bold"],
                "italic": font.italic if action.get("italic") is None else action["italic"],
                "underline": font.underline,
                "strike": font.strike,
                "color": font.color,
            }
            if action.get("font_color"):
                kwargs["color"] = action["font_color"].lstrip("#").upper()
            cell.font = Font(**kwargs)
            if action.get("number_format"):
                cell.number_format = action["number_format"]
        if action.get("background"):
            self._highlight_cells(parsed, action["background"])

    def _do_set_column_width(self, action: dict[str, Any]) -> None:
        self._ws(action["sheet"]).column_dimensions[str(action["column"]).upper()].width = float(action["width"])

    def _do_freeze_panes(self, action: dict[str, Any]) -> None:
        self._ws(action["sheet"]).freeze_panes = str(action["cell"]).upper()

    def _do_add_note(self, action: dict[str, Any]) -> None:
        parsed = parse_range(action["cell"])
        cell = self._ws(parsed.sheet).cell(row=parsed.start_row + 1, column=parsed.start_col + 1)
        cell.comment = Comment(str(action["text"]), "Vitreus")

    def _do_sort_range(self, action: dict[str, Any]) -> None:
        parsed = parse_range(action["range"])
        ws = self._ws(parsed.sheet)
        key_col = column_index(str(action["by_column"]))
        if not (parsed.start_col <= key_col <= parsed.end_col):
            raise ValueError(f"by_column {action['by_column']} is outside {action['range']}")
        first_data_row = parsed.start_row + (1 if action.get("has_header", True) else 0)
        block = [
            [ws.cell(row=r + 1, column=c + 1).value for c in range(parsed.start_col, parsed.end_col + 1)]
            for r in range(first_data_row, parsed.end_row + 1)
        ]
        rel = key_col - parsed.start_col
        block.sort(key=lambda row: _sort_key(row[rel]), reverse=bool(action.get("descending", False)))
        for r_offset, row in enumerate(block):
            for c_offset, value in enumerate(row):
                ws.cell(row=first_data_row + 1 + r_offset, column=parsed.start_col + 1 + c_offset).value = value

    def _do_insert_rows(self, action: dict[str, Any]) -> None:
        self._ws(action["sheet"]).insert_rows(int(action["at"]), int(action.get("count", 1)))

    def _do_delete_rows(self, action: dict[str, Any]) -> None:
        self._ws(action["sheet"]).delete_rows(int(action["at"]), int(action.get("count", 1)))

    def _do_clear_range(self, action: dict[str, Any]) -> None:
        parsed = parse_range(action["range"])
        for r, c, cell in self._cells(parsed):
            cell.value = None
            self._formats.pop(f"{parsed.sheet}!{column_name(c)}{r + 1}", None)

    def _do_add_sheet(self, action: dict[str, Any]) -> None:
        name = str(action["name"])
        if name in self.workbook.sheetnames:
            raise ValueError(f"Sheet already exists: {name}")
        self.workbook.create_sheet(title=name)

    def _do_add_chart(self, action: dict[str, Any]) -> None:
        ws = self._ws(action["sheet"])
        data = parse_range(action["data_range"])
        data_ws = self._ws(data.sheet)
        chart_type = action.get("chart_type", "bar")
        min_col, max_col = data.start_col + 1, data.end_col + 1
        min_row, max_row = data.start_row + 1, data.end_row + 1
        if chart_type == "scatter":
            chart: Any = ScatterChart()
            x_values = Reference(data_ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
            for col in range(min_col + 1, max_col + 1):
                y_values = Reference(data_ws, min_col=col, min_row=min_row, max_row=max_row)
                chart.series.append(Series(y_values, x_values, title_from_data=True))
        else:
            chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[chart_type]()
            values = Reference(data_ws, min_col=min_col + 1 if max_col > min_col else min_col, max_col=max_col, min_row=min_row, max_row=max_row)
            chart.add_data(values, titles_from_data=True)
            if max_col > min_col:
                chart.set_categories(Reference(data_ws, min_col=min_col, min_row=min_row + 1, max_row=max_row))
        if action.get("title"):
            chart.title = action["title"]
        ws.add_chart(chart, str(action.get("anchor", "H2")).upper())

    # ── persistence ──

    def save(self, path: str | Path) -> list[Path]:
        path = Path(path)
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            self.workbook.save(path)
            return [path]
        snapshot = self.snapshot()
        snapshot.save_csv(str(path), sheet_name=self.active_sheet)
        written = [path]
        if self._formats:
            sidecar = path.with_name(f"{path.stem}_highlights.json")
            sidecar.write_text(
                json.dumps({ref: {"background": fmt.background} for ref, fmt in self._formats.items()}, indent=2),
                encoding="utf-8",
            )
            written.append(sidecar)
        return written


def _sort_key(value: Any) -> tuple[int, Any]:
    if value is None or value == "":
        return (2, "")
    if isinstance(value, (int, float)):
        return (0, value)
    return (1, str(value).lower())


InMemoryCalcDriver = WorkbookDriver
