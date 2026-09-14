import shutil
import time
from pathlib import Path

import openpyxl
import pytest

from core.uno_driver import UnoDriver, find_uno_python, launch_calc

pytestmark = pytest.mark.integration


def _terminate(process):
    process.terminate()
    try:
        process.wait(timeout=10)
    except Exception:
        process.kill()
        process.wait(timeout=10)


def test_uno_driver_applies_manifest_to_live_calc_document(tmp_path: Path):
    if not shutil.which("soffice") or not find_uno_python():
        pytest.skip("LibreOffice soffice and PyUNO-capable Python are required")

    source = Path("examples/test_workbook.xlsx")
    workbook = tmp_path / "test_workbook.xlsx"
    shutil.copy2(source, workbook)
    profile_dir = tmp_path / "lo-profile"
    port = 2199

    soffice = launch_calc(str(workbook), port=port, headless=True, user_profile=str(profile_dir))
    driver = UnoDriver(port=port, open_path=None, timeout=10)
    try:
        deadline = time.monotonic() + 40
        while time.monotonic() < deadline:
            if driver.is_available():
                break
            time.sleep(1)
        else:
            pytest.fail("UNO bridge did not become available within 40 seconds")

        snapshot = driver.snapshot()
        assert {"Sales", "Expenses", "HR_Reviews"}.issubset(snapshot.sheets)

        summary = driver.execute_manifest(
            {
                "actions": [
                    {"type": "highlight", "range": "Expenses!A2:D2", "color": "#ef4444", "reason": "integration marker"},
                    {"type": "write_value", "cell": "Expenses!Z1", "value": "VITREUS"},
                    {"type": "formula", "cell": "Expenses!Z2", "formula": "=1+1"},
                    {"type": "add_note", "cell": "Expenses!A2", "text": "Reviewed by Vitreus"},
                    {"type": "set_format", "range": "Expenses!A1:D1", "bold": True},
                    {"type": "add_sheet", "name": "Summary"},
                    {"type": "write_value", "cell": "Summary!A1", "value": 42},
                    {
                        "type": "add_chart",
                        "sheet": "Expenses",
                        "chart_type": "bar",
                        "data_range": "Expenses!A1:B5",
                        "title": "Expenses",
                        "anchor": "H2",
                    },
                ]
            }
        )
        assert summary.applied == 8
        assert summary.errors == []

        out = tmp_path / "out.xlsx"
        assert driver.save(str(out)) == str(out)

        saved = openpyxl.load_workbook(out)
        assert saved["Expenses"]["Z1"].value == "VITREUS"
        assert saved["Expenses"]["A2"].fill.fgColor.rgb.endswith("EF4444")
        assert saved["Expenses"]["A1"].font.bold is True
        assert "Summary" in saved.sheetnames
        assert saved["Summary"]["A1"].value == 42
    finally:
        driver.close()
        _terminate(soffice)
