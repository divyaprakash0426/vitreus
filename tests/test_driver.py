import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from core.driver import (
    CellChange,
    CellFormat,
    InMemoryCalcDriver,
    WorkbookDriver,
    WorkbookSnapshot,
    diff_snapshots,
)
from core.manifest import validate_manifest


def _driver(rows, sheet="Sheet1") -> WorkbookDriver:
    return WorkbookDriver.from_snapshot(WorkbookSnapshot(sheets={sheet: rows}))


SCORES = [["Name", "Score"], ["Ada", 91], ["Linus", 72]]


# ─── WorkbookSnapshot ────────────────────────────────────────────────────────


def test_workbook_snapshot_exports_range_to_csv_and_json():
    snapshot = WorkbookSnapshot(sheets={"Budget": [["Category", "Amount"], ["Cloud", 120.5], ["Hardware", 240]]})

    assert snapshot.range_to_csv("Budget!A1:B3") == "Category,Amount\r\nCloud,120.5\r\nHardware,240\r\n"
    assert json.loads(snapshot.range_to_json("Budget!A2:B3")) == [
        {"Category": "Cloud", "Amount": 120.5},
        {"Category": "Hardware", "Amount": 240},
    ]


def test_workbook_snapshot_dims_and_data_range():
    snapshot = WorkbookSnapshot(sheets={"S": [["a", "b", "c"], [1, 2, 3], [4, 5, 6], [7, 8, 9]]})

    assert snapshot.dims("S") == (4, 3)
    assert snapshot.data_range("S") == "S!A1:C4"
    assert WorkbookSnapshot(sheets={"Empty": []}).data_range("Empty") == "Empty!A1:A1"


def test_workbook_snapshot_from_csv_text_coerces_numbers():
    snapshot = WorkbookSnapshot.from_csv_text("Name,Score\nAda,91\nLinus,72.5\n", sheet_name="Data")

    assert snapshot.sheets["Data"] == [["Name", "Score"], ["Ada", 91], ["Linus", 72.5]]


def test_workbook_snapshot_saves_sheet_to_csv_file(tmp_path: Path):
    snapshot = WorkbookSnapshot(sheets={"Sheet1": [["Name", "Score", "Status"], ["Ada", 91, "pass"]]})

    out = tmp_path / "out.csv"
    snapshot.save_csv(str(out), sheet_name="Sheet1")

    assert out.read_text(encoding="utf-8").splitlines() == ["Name,Score,Status", "Ada,91,pass"]


# ─── WorkbookDriver: core actions ────────────────────────────────────────────


def test_driver_executes_highlight_value_and_formula_manifest():
    driver = _driver([row[:] for row in SCORES])

    summary = driver.execute_manifest(
        {
            "actions": [
                {"type": "highlight", "range": "Sheet1!B2:B2", "color": "#16a34a"},
                {"type": "write_value", "cell": "Sheet1!C2", "value": "pass"},
                {"type": "formula", "cell": "Sheet1!C3", "formula": '=IF(B3>80,"pass","review")'},
            ]
        }
    )

    assert summary.applied == 3 and summary.errors == []
    sheet = driver.snapshot().sheets["Sheet1"]
    assert sheet[1][2] == "pass"
    assert sheet[2][2] == '=IF(B3>80,"pass","review")'
    assert driver.formats["Sheet1!B2"] == CellFormat(background="#16a34a")


def test_driver_accepts_validated_manifest_object():
    driver = _driver([row[:] for row in SCORES])
    manifest = validate_manifest({"actions": [{"type": "write_value", "cell": "Sheet1!A4", "value": "x"}]}, {"Sheet1"})

    assert driver.execute_manifest(manifest).applied == 1
    assert driver.snapshot().sheets["Sheet1"][3][0] == "x"


def test_in_memory_alias_accepts_snapshot_positional_argument():
    driver = InMemoryCalcDriver(WorkbookSnapshot(sheets={"Sheet1": [["A"], [1]]}))

    summary = driver.execute_manifest({"actions": [{"type": "delete_sheet", "sheet": "Sheet1"}]})

    assert summary.applied == 0
    assert summary.errors == ["Unsupported action type: delete_sheet"]
    assert driver.snapshot().sheets["Sheet1"] == [["A"], [1]]


def test_write_range_fills_a_block_and_grows_the_sheet():
    driver = _driver([["A"], [1]])

    driver.execute_manifest({"actions": [{"type": "write_range", "range": "Sheet1!B1:C2", "values": [["x", "y"], [2, 3]]}]})

    assert driver.snapshot().sheets["Sheet1"] == [["A", "x", "y"], [1, 2, 3]]


def test_fill_formula_with_row_placeholder():
    driver = _driver([["A", "B"], [1, 2], [3, 4], [5, 6]])

    driver.execute_manifest({"actions": [{"type": "fill_formula", "range": "Sheet1!C2:C4", "formula": "=A{row}+B{row}"}]})

    col = [row[2] for row in driver.snapshot().sheets["Sheet1"][1:]]
    assert col == ["=A2+B2", "=A3+B3", "=A4+B4"]


def test_fill_formula_shifts_relative_references_but_not_absolute():
    driver = _driver([["A", "B"], [1, 2], [3, 4], [5, 6]])

    driver.execute_manifest({"actions": [{"type": "fill_formula", "range": "Sheet1!C2:C4", "formula": "=A2*$B$2"}]})

    col = [row[2] for row in driver.snapshot().sheets["Sheet1"][1:]]
    assert col == ["=A2*$B$2", "=A3*$B$2", "=A4*$B$2"]


def test_set_format_applies_font_fill_and_number_format():
    driver = _driver([["Name", "Amount"], ["Ada", 1234.5]])

    driver.execute_manifest(
        {
            "actions": [
                {"type": "set_format", "range": "Sheet1!A1:B1", "bold": True, "background": "#e5e7eb", "font_color": "#111111"},
                {"type": "set_format", "range": "Sheet1!B2", "number_format": "#,##0.00", "italic": True},
            ]
        }
    )

    ws = driver.workbook["Sheet1"]
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.fgColor.rgb.endswith("E5E7EB")
    assert ws["A1"].font.color.rgb.endswith("111111")
    assert ws["B2"].number_format == "#,##0.00"
    assert ws["B2"].font.italic is True
    assert driver.formats["Sheet1!B1"].background == "#e5e7eb"


def test_add_sheet_then_write_into_it():
    driver = _driver([["A"]])

    summary = driver.execute_manifest(
        {"actions": [{"type": "add_sheet", "name": "Summary"}, {"type": "write_value", "cell": "Summary!A1", "value": 42}]}
    )

    assert summary.applied == 2
    assert driver.snapshot().sheets["Summary"] == [[42]]


def test_insert_and_delete_rows_shift_data():
    driver = _driver([["H"], [1], [2], [3]])

    driver.execute_manifest({"actions": [{"type": "insert_rows", "sheet": "Sheet1", "at": 2, "count": 1}]})
    assert driver.snapshot().sheets["Sheet1"] == [["H"], [""], [1], [2], [3]]

    driver.execute_manifest({"actions": [{"type": "delete_rows", "sheet": "Sheet1", "at": 2, "count": 2}]})
    assert driver.snapshot().sheets["Sheet1"] == [["H"], [2], [3]]


def test_sort_range_descending_keeps_header_in_place():
    driver = _driver([["Name", "Score"], ["Ada", 91], ["Linus", 72], ["Grace", 88]])

    driver.execute_manifest(
        {"actions": [{"type": "sort_range", "range": "Sheet1!A1:B4", "by_column": "B", "descending": True, "has_header": True}]}
    )

    assert driver.snapshot().sheets["Sheet1"] == [["Name", "Score"], ["Ada", 91], ["Grace", 88], ["Linus", 72]]


def test_sort_range_by_text_column_ascending_without_header():
    driver = _driver([["b", 2], ["a", 1], ["c", 3]])

    driver.execute_manifest(
        {"actions": [{"type": "sort_range", "range": "Sheet1!A1:B3", "by_column": "A", "has_header": False}]}
    )

    assert driver.snapshot().sheets["Sheet1"] == [["a", 1], ["b", 2], ["c", 3]]


def test_clear_range_empties_cells():
    driver = _driver([["A", "B"], [1, 2]])

    driver.execute_manifest({"actions": [{"type": "clear_range", "range": "Sheet1!B1:B2"}]})

    assert driver.snapshot().sheets["Sheet1"] == [["A", ""], [1, ""]]


def test_add_note_column_width_freeze_panes_and_chart():
    driver = _driver([["Name", "Score"], ["Ada", 91], ["Linus", 72]])

    summary = driver.execute_manifest(
        {
            "actions": [
                {"type": "add_note", "cell": "Sheet1!A2", "text": "Reviewed by Vitreus"},
                {"type": "set_column_width", "sheet": "Sheet1", "column": "A", "width": 30},
                {"type": "freeze_panes", "sheet": "Sheet1", "cell": "A2"},
                {"type": "add_chart", "sheet": "Sheet1", "chart_type": "bar", "data_range": "Sheet1!A1:B3", "title": "Scores", "anchor": "D2"},
            ]
        }
    )

    ws = driver.workbook["Sheet1"]
    assert summary.applied == 4 and summary.errors == []
    assert ws["A2"].comment.text == "Reviewed by Vitreus"
    assert ws.column_dimensions["A"].width == 30
    assert ws.freeze_panes == "A2"
    assert len(ws._charts) == 1


def test_chart_type_variants_all_apply():
    for chart_type in ("bar", "line", "pie", "scatter"):
        driver = _driver([["Name", "Score"], ["Ada", 91], ["Linus", 72]])
        summary = driver.execute_manifest(
            {"actions": [{"type": "add_chart", "sheet": "Sheet1", "chart_type": chart_type, "data_range": "Sheet1!A1:B3"}]}
        )
        assert summary.errors == [], chart_type


def test_action_errors_are_reported_and_do_not_stop_later_actions():
    driver = _driver([["A"], [1]])

    summary = driver.execute_manifest(
        {
            "actions": [
                {"type": "write_value", "cell": "Ghost!A1", "value": 1},
                {"type": "write_value", "cell": "Sheet1!A3", "value": "kept"},
            ]
        }
    )

    assert summary.applied == 1
    assert len(summary.errors) == 1 and "Ghost" in summary.errors[0]
    assert driver.snapshot().sheets["Sheet1"][2][0] == "kept"


# ─── Loading and saving ──────────────────────────────────────────────────────


def test_from_path_csv_names_sheet_after_option_or_default(tmp_path: Path):
    csv_path = tmp_path / "budget.csv"
    csv_path.write_text("Name,Score\nAda,91\n", encoding="utf-8")

    assert WorkbookDriver.from_path(csv_path).active_sheet == "Sheet1"
    assert WorkbookDriver.from_path(csv_path, sheet_name="Data").snapshot().sheets["Data"][1] == ["Ada", 91]


def test_xlsx_round_trip_preserves_other_sheets_and_existing_styles(tmp_path: Path):
    src = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Score"])
    ws.append(["Ada", 91])
    ws["A1"].font = Font(bold=True)
    other = wb.create_sheet("Notes")
    other["A1"] = "keep me"
    wb.save(src)

    driver = WorkbookDriver.from_path(src, sheet_name="Data")
    driver.execute_manifest({"actions": [{"type": "highlight", "range": "Data!A2:B2", "color": "#f97316"}]})
    out = tmp_path / "out.xlsx"
    saved = driver.save(out)

    assert saved == [out]
    reloaded = openpyxl.load_workbook(out)
    assert reloaded.sheetnames == ["Data", "Notes"]
    assert reloaded["Notes"]["A1"].value == "keep me"
    assert reloaded["Data"]["A1"].font.bold is True
    assert reloaded["Data"]["A2"].fill.fgColor.rgb.endswith("F97316")


def test_save_csv_writes_values_and_highlight_sidecar(tmp_path: Path):
    driver = _driver([["Name", "Score"], ["Ada", 91]])
    driver.execute_manifest(
        {
            "actions": [
                {"type": "highlight", "range": "Sheet1!A2:B2", "color": "#f97316"},
                {"type": "write_value", "cell": "Sheet1!C2", "value": "ok"},
            ]
        }
    )

    out = tmp_path / "result.csv"
    saved = driver.save(out)

    sidecar = tmp_path / "result_highlights.json"
    assert saved == [out, sidecar]
    assert out.read_text(encoding="utf-8").splitlines()[1] == "Ada,91,ok"
    assert json.loads(sidecar.read_text())["Sheet1!A2"] == {"background": "#f97316"}


def test_save_csv_without_highlights_has_no_sidecar(tmp_path: Path):
    driver = _driver([["A"], [1]])
    driver.execute_manifest({"actions": [{"type": "write_value", "cell": "Sheet1!A2", "value": 2}]})

    out = tmp_path / "plain.csv"

    assert driver.save(out) == [out]
    assert not (tmp_path / "plain_highlights.json").exists()


def test_from_csv_text_and_source_name():
    driver = WorkbookDriver.from_csv_text("a,b\n1,2\n", sheet_name="Piped")

    assert driver.snapshot().sheets["Piped"] == [["a", "b"], [1, 2]]
    assert driver.source == "stdin"


# ─── Diff ────────────────────────────────────────────────────────────────────


def test_diff_snapshots_reports_only_changed_cells():
    before = WorkbookSnapshot(sheets={"S": [["A", "B"], [1, 2]]})
    after = WorkbookSnapshot(sheets={"S": [["A", "B"], [1, 5], ["", "new"]], "T": [["t"]]})

    changes = diff_snapshots(before, after)

    assert changes == [
        CellChange(sheet="S", cell="B2", old=2, new=5),
        CellChange(sheet="S", cell="B3", old=None, new="new"),
        CellChange(sheet="T", cell="A1", old=None, new="t"),
    ]


def test_diff_snapshots_is_empty_for_identical_snapshots():
    snap = WorkbookSnapshot(sheets={"S": [["A"], [1]]})

    assert diff_snapshots(snap, snap) == []


def test_csv_round_trip_preserves_ids_zip_codes_and_exponent_strings(tmp_path: Path):
    text = "ID,Zip,Phone,Exp,Amount\n007,02134,+91 98,1e3,1.50\n"
    driver = WorkbookDriver.from_csv_text(text)

    out = tmp_path / "out.csv"
    driver.save(out)

    # Plain decimals normalise (1.50 → 1.5) exactly as a spreadsheet import would; codes stay text.
    assert out.read_text(encoding="utf-8").splitlines()[1] == "007,02134,+91 98,1e3,1.5"


def test_parse_range_accepts_quoted_sheets_and_lowercase_cells():
    from core.driver import parse_range

    parsed = parse_range("'My Sheet'!a2:c3")

    assert (parsed.sheet, parsed.start_col, parsed.start_row, parsed.end_col, parsed.end_row) == ("My Sheet", 0, 1, 2, 2)


def test_delete_rows_rewrites_dependent_formulas():
    driver = WorkbookDriver.from_csv_text("V\n1\n2\n3\n4\n")
    ws = driver.workbook["Sheet1"]
    ws["B6"] = "=SUM(A2:A5)"
    ws["C6"] = "=A5*2"
    ws["D6"] = '=IF(A2>0,"See A2 above","")'
    driver.workbook.create_sheet("Other")["A1"] = "=Sheet1!A4+'Sheet1'!$A$5"

    driver.execute_manifest({"actions": [{"type": "delete_rows", "sheet": "Sheet1", "at": 2, "count": 1}]})

    assert ws["B5"].value == "=SUM(A2:A4)"
    assert ws["C5"].value == "=A4*2"
    assert ws["D5"].value == '=IF(#REF!>0,"See A2 above","")'
    assert driver.workbook["Other"]["A1"].value == "=Sheet1!A3+'Sheet1'!$A$4"


def test_insert_rows_shifts_formulas_and_highlight_formats():
    driver = WorkbookDriver.from_csv_text("V\n1\n2\n")
    ws = driver.workbook["Sheet1"]
    ws["B4"] = "=SUM(A2:A3)"
    driver.execute_manifest({"actions": [{"type": "highlight", "range": "Sheet1!A3:A3", "color": "#f97316", "reason": "r"}]})

    driver.execute_manifest({"actions": [{"type": "insert_rows", "sheet": "Sheet1", "at": 2, "count": 2}]})

    assert ws["B6"].value == "=SUM(A4:A5)"
    assert "Sheet1!A5" in driver.formats and "Sheet1!A3" not in driver.formats


def test_xlsx_snapshot_prefers_cached_values_over_formula_text(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["A", "Total"])
    ws.append([2, "=A2*10"])
    path = tmp_path / "book.xlsx"
    wb.save(path)
    # Simulate a cached value as Excel/LibreOffice would store it.
    import zipfile, shutil
    tmp = tmp_path / "book2.xlsx"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>A2*10</f><v></v>", b"<f>A2*10</f><v>20</v>").replace(b"<f>A2*10</f>", b"<f>A2*10</f><v>20</v>", 1) if b"<v>20</v>" not in data else data
            zout.writestr(item, data)
    shutil.move(tmp, path)

    snapshot = WorkbookDriver.from_path(path).snapshot()

    assert snapshot.sheets["Data"][1] == [2, 20]


def test_coerce_value_accepts_plain_decimals_but_not_ids_or_nan():
    from core.driver import coerce_value

    assert coerce_value("12.50") == 12.5 and coerce_value("100.00") == 100.0 and coerce_value("-0.5") == -0.5
    assert coerce_value("007") == "007" and coerce_value("+91") == "+91" and coerce_value("1e3") == "1e3"
    assert coerce_value("nan") == "nan" and coerce_value("inf") == "inf" and coerce_value("01.5") == "01.5"


def test_sort_range_orders_numeric_text_numerically():
    driver = WorkbookDriver.from_csv_text('Item,Amount\na,12.50\nb,"1,234.00"\nc,7.25\nd,50\n')

    driver.execute_manifest({"actions": [{"type": "sort_range", "range": "Sheet1!A1:B5", "by_column": "B", "descending": True}]})

    assert [r[1] for r in driver.snapshot().sheets["Sheet1"][1:]] == ["1,234.00", 50, 12.5, 7.25]


def test_row_edit_rewrites_lowercase_refs():
    from core.driver import adjust_formula_for_row_edit

    assert adjust_formula_for_row_edit("=sum(a5:a9)+b10", "S", "S", 2, 1, True) == "=sum(a4:a8)+b9"
