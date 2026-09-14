import json
from pathlib import Path

import openpyxl
from typer.testing import CliRunner

from interfaces.cli import app

runner = CliRunner()

MOCK_MANIFEST = json.dumps(
    {
        "summary": "Linus needs review.",
        "actions": [{"type": "highlight", "range": "Sheet1!A3:B3", "color": "#f97316", "reason": "low score"}],
    }
)
CSV = "Name,Score\nAda,91\nLinus,72\n"


def mock_google(monkeypatch, reply: str = MOCK_MANIFEST):
    calls: list[list[dict]] = []

    def chat(self, messages, images=None):
        calls.append([dict(m) for m in messages])
        return reply

    monkeypatch.setattr("core.backends.GoogleAIBackend.chat", chat)
    return calls


# ─── models / doctor / config ───────────────────────────────────────────────


def test_models_command_shows_gemma_4_policy():
    result = runner.invoke(app, ["models"])

    assert result.exit_code == 0, result.output
    payload = json.loads(result.stdout)
    assert payload["policy"]["primary"] == "gemma4:31b"
    assert payload["policy"]["drafter"] == "gemma4:e4b"
    assert "31B Dense" in payload["policy"]["rationale"]
    assert payload["backend"] == "fallback"
    assert "gemma4:31b" in result.stdout


def test_models_lists_backend_models(monkeypatch):
    monkeypatch.setattr("core.backends.GoogleAIBackend.list_models", lambda self: ["gemma-4-31b-it", "gemma-4-26b-a4b-it"])

    result = runner.invoke(app, ["models", "--backend", "google", "--api-key", "k"])

    assert result.exit_code == 0, result.output
    payload = json.loads(result.stdout)
    assert payload["backend"] == "google" and payload["model"] == "gemma-4-31b-it"
    assert payload["available"] == ["gemma-4-31b-it", "gemma-4-26b-a4b-it"]


def test_doctor_reports_environment_as_json(monkeypatch):
    monkeypatch.setattr("interfaces.cli.ollama_reachable", lambda host, **kw: False)

    result = runner.invoke(app, ["doctor"])

    assert result.exit_code == 0, result.output
    report = json.loads(result.stdout)
    assert report["ollama"]["reachable"] is False
    assert report["api_keys"] == {"gemini": False, "openrouter": False, "openai": False}
    assert report["selected_backend"]["backend"] == "fallback"
    assert "libreoffice" in report and "openpyxl" in report


def test_config_init_writes_template_and_show_prints_effective_settings(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("XDG_CONFIG_HOME", str(tmp_path))

    init = runner.invoke(app, ["config", "--init"])
    assert init.exit_code == 0, init.output
    config_path = tmp_path / "vitreus" / "config.toml"
    assert config_path.exists() and 'backend = "auto"' in config_path.read_text()

    again = runner.invoke(app, ["config", "--init"])
    assert again.exit_code != 0  # refuses to overwrite

    show = runner.invoke(app, ["config", "--show"])
    assert show.exit_code == 0
    assert json.loads(show.stdout)["primary"] == "gemma4:31b"


# ─── analyze ────────────────────────────────────────────────────────────────


def test_analyze_fallback_prints_manifest_and_notice(tmp_path: Path):
    csv_path = tmp_path / "scores.csv"
    csv_path.write_text(CSV, encoding="utf-8")

    result = runner.invoke(app, ["analyze", str(csv_path), "Highlight rows that need review"])

    assert result.exit_code == 0, result.output
    manifest = json.loads(result.stdout)
    assert manifest["actions"][0]["range"] == "Sheet1!A3:B3"
    assert manifest["model"]["backend"] == "fallback"
    assert "fallback" in result.stderr.lower()


def test_analyze_reads_csv_from_stdin():
    result = runner.invoke(app, ["analyze", "-", "Highlight rows that need review"], input=CSV)

    assert result.exit_code == 0, result.output
    assert json.loads(result.stdout)["actions"][0]["range"] == "Sheet1!A3:B3"


def test_analyze_with_google_backend_uses_mocked_chat(tmp_path: Path, monkeypatch):
    csv_path = tmp_path / "scores.csv"
    csv_path.write_text(CSV, encoding="utf-8")
    calls = mock_google(monkeypatch)

    result = runner.invoke(app, ["analyze", str(csv_path), "any query", "--backend", "google", "--api-key", "dummy"])

    assert result.exit_code == 0, result.output
    manifest = json.loads(result.stdout)
    assert manifest["model"]["backend"] == "google" and manifest["model"]["model"] == "gemma-4-31b-it"
    assert manifest["summary"] == "Linus needs review."
    assert "Linus" in calls[0][-1]["content"]


def test_analyze_reads_api_key_from_env(tmp_path: Path, monkeypatch):
    csv_path = tmp_path / "scores.csv"
    csv_path.write_text(CSV, encoding="utf-8")
    mock_google(monkeypatch)

    result = runner.invoke(app, ["analyze", str(csv_path), "q", "--backend", "google"], env={"GEMINI_API_KEY": "env-key"})

    assert result.exit_code == 0, result.output


def test_analyze_google_without_key_fails_with_hint(tmp_path: Path):
    csv_path = tmp_path / "scores.csv"
    csv_path.write_text(CSV, encoding="utf-8")

    result = runner.invoke(app, ["analyze", str(csv_path), "q", "--backend", "google"])

    assert result.exit_code == 1
    assert "GEMINI_API_KEY" in result.stderr


def test_analyze_fast_flag_selects_drafter_model(tmp_path: Path, monkeypatch):
    csv_path = tmp_path / "scores.csv"
    csv_path.write_text(CSV, encoding="utf-8")
    mock_google(monkeypatch)

    result = runner.invoke(app, ["analyze", str(csv_path), "q", "--backend", "google", "--api-key", "k", "--fast"])

    assert json.loads(result.stdout)["model"]["model"] == "gemma-4-26b-a4b-it"


def test_analyze_preview_prints_diff_to_stderr_without_writing(tmp_path: Path, monkeypatch):
    csv_path = tmp_path / "scores.csv"
    csv_path.write_text(CSV, encoding="utf-8")
    out = tmp_path / "out.xlsx"
    reply = json.dumps({"summary": "s", "actions": [{"type": "write_value", "cell": "Sheet1!C2", "value": "review"}]})
    mock_google(monkeypatch, reply)

    result = runner.invoke(app, ["analyze", str(csv_path), "q", "--backend", "google", "--api-key", "k", "--preview", "--output", str(out)])

    assert result.exit_code == 0, result.output
    assert "C2" in result.stderr and "review" in result.stderr
    assert not out.exists()
    assert json.loads(result.stdout)["actions"][0]["cell"] == "Sheet1!C2"


def test_analyze_in_place_overwrites_source(tmp_path: Path, monkeypatch):
    xlsx = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.active.append(["Name", "Score"])
    wb.active.append(["Ada", 91])
    wb.save(xlsx)
    reply = json.dumps({"actions": [{"type": "write_value", "cell": "Sheet1!C2", "value": "ok"}]})
    mock_google(monkeypatch, reply)

    result = runner.invoke(app, ["analyze", str(xlsx), "q", "--backend", "google", "--api-key", "k", "--in-place"])

    assert result.exit_code == 0, result.output
    assert json.loads(result.stdout)["saved"] == str(xlsx)
    assert openpyxl.load_workbook(xlsx)["Sheet1"]["C2"].value == "ok"


def test_analyze_all_sheets_and_image_options(tmp_path: Path, monkeypatch):
    from PIL import Image

    xlsx = tmp_path / "wb.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Sales"
    wb.active.append(["Item"])
    wb.create_sheet("Costs").append(["Cost"])
    wb.save(xlsx)
    img = tmp_path / "r.png"
    Image.new("RGB", (8, 8)).save(img)
    seen = {}

    def chat(self, messages, images=None):
        seen["images"] = images
        seen["content"] = messages[-1]["content"]
        return json.dumps({"actions": []})

    monkeypatch.setattr("core.backends.GoogleAIBackend.chat", chat)

    result = runner.invoke(app, ["analyze", str(xlsx), "q", "--backend", "google", "--api-key", "k", "--all-sheets", "--image", str(img)])

    assert result.exit_code == 0, result.output
    assert len(seen["images"]) == 1
    assert "Costs" in seen["content"] and "Cost" in seen["content"]


def test_analyze_reports_agent_failure_cleanly(tmp_path: Path, monkeypatch):
    csv_path = tmp_path / "scores.csv"
    csv_path.write_text(CSV, encoding="utf-8")
    mock_google(monkeypatch, "I cannot help with that.")

    result = runner.invoke(app, ["analyze", str(csv_path), "q", "--backend", "google", "--api-key", "k"])

    assert result.exit_code == 1
    assert "valid manifest" in result.stderr


# ─── ask ────────────────────────────────────────────────────────────────────


def test_ask_prints_summary_only(tmp_path: Path, monkeypatch):
    csv_path = tmp_path / "scores.csv"
    csv_path.write_text(CSV, encoding="utf-8")
    mock_google(monkeypatch, json.dumps({"summary": "Ada leads with 91.", "actions": []}))

    result = runner.invoke(app, ["ask", str(csv_path), "Who leads?", "--backend", "google", "--api-key", "k"])

    assert result.exit_code == 0, result.output
    assert result.stdout.strip() == "Ada leads with 91."


# ─── apply-manifest / batch ─────────────────────────────────────────────────


def test_apply_manifest_updates_csv_snapshot(tmp_path: Path):
    csv_path = tmp_path / "sheet.csv"
    manifest_path = tmp_path / "manifest.json"
    csv_path.write_text("Name,Score\nAda,91\n", encoding="utf-8")
    manifest_path.write_text(json.dumps({"actions": [{"type": "write_value", "cell": "Sheet1!C2", "value": "pass"}]}), encoding="utf-8")

    result = runner.invoke(app, ["apply-manifest", str(csv_path), str(manifest_path)])

    assert result.exit_code == 0, result.output
    assert json.loads(result.stdout) == {"applied": 1, "errors": []}


def test_apply_manifest_with_output_saves_xlsx(tmp_path: Path):
    csv_path = tmp_path / "sheet.csv"
    manifest_path = tmp_path / "manifest.json"
    out = tmp_path / "out.xlsx"
    csv_path.write_text("Name,Score\nAda,91\n", encoding="utf-8")
    manifest_path.write_text(json.dumps({"actions": [{"type": "formula", "cell": "Sheet1!C2", "formula": "=B2*2"}]}), encoding="utf-8")

    result = runner.invoke(app, ["apply-manifest", str(csv_path), str(manifest_path), "--output", str(out)])

    assert result.exit_code == 0, result.output
    assert json.loads(result.stdout)["saved"] == str(out)
    assert openpyxl.load_workbook(out)["Sheet1"]["C2"].value == "=B2*2"


def test_apply_manifest_rejects_invalid_manifest(tmp_path: Path):
    csv_path = tmp_path / "sheet.csv"
    manifest_path = tmp_path / "manifest.json"
    csv_path.write_text("Name,Score\nAda,91\n", encoding="utf-8")
    manifest_path.write_text(json.dumps({"actions": [{"type": "nuke"}]}), encoding="utf-8")

    result = runner.invoke(app, ["apply-manifest", str(csv_path), str(manifest_path)])

    assert result.exit_code == 1
    assert "Unsupported action type" in result.stderr


def test_batch_processes_each_file_into_output_dir(tmp_path: Path):
    for name in ("a", "b"):
        (tmp_path / f"{name}.csv").write_text(CSV, encoding="utf-8")
    out_dir = tmp_path / "out"

    result = runner.invoke(app, ["batch", "Highlight rows that need review", str(tmp_path / "a.csv"), str(tmp_path / "b.csv"), "--output-dir", str(out_dir)])

    assert result.exit_code == 0, result.output
    lines = [json.loads(line) for line in result.stdout.strip().splitlines()]
    assert [Path(entry["file"]).name for entry in lines] == ["a.csv", "b.csv"]
    assert all(entry["applied"] == 1 for entry in lines)
    assert (out_dir / "a.xlsx").exists() and (out_dir / "b.xlsx").exists()
    fill = openpyxl.load_workbook(out_dir / "a.xlsx")["Sheet1"]["A3"].fill
    assert fill.fgColor.rgb.upper().endswith("F97316")


def test_batch_continues_after_a_bad_file_and_exits_nonzero(tmp_path: Path):
    good = tmp_path / "good.csv"
    good.write_text(CSV, encoding="utf-8")
    out_dir = tmp_path / "out"

    result = runner.invoke(app, ["batch", "Highlight rows that need review", str(tmp_path / "missing.csv"), str(good), "--output-dir", str(out_dir)])

    assert result.exit_code == 1
    lines = [json.loads(line) for line in result.stdout.strip().splitlines()]
    assert "error" in lines[0] and lines[1]["applied"] == 1


def test_batch_rejects_file_path_given_as_instruction(tmp_path: Path):
    good = tmp_path / "good.csv"
    good.write_text(CSV, encoding="utf-8")

    result = runner.invoke(app, ["batch", str(good), "Highlight rows that need review", "--output-dir", str(tmp_path / "out")])

    assert result.exit_code == 1
    assert "first argument must be the instruction" in result.stderr.lower()
    assert 'vitreus batch "instruction"' in result.stderr.lower()
    assert not (tmp_path / "out").exists()


# ─── vision ─────────────────────────────────────────────────────────────────


def test_vision_without_backend_prints_metadata_payload(tmp_path: Path):
    from PIL import Image

    img = tmp_path / "chart.png"
    Image.new("RGB", (32, 16)).save(img)

    result = runner.invoke(app, ["vision", str(img), "--purpose", "chart"])

    assert result.exit_code == 0, result.output
    payload = json.loads(result.stdout)
    assert payload["purpose"] == "chart" and payload["width"] == 32


def test_vision_with_backend_extracts_rows_and_saves_workbook(tmp_path: Path, monkeypatch):
    from PIL import Image

    img = tmp_path / "receipt.png"
    Image.new("RGB", (32, 16)).save(img)
    out = tmp_path / "receipt.xlsx"
    mock_google(monkeypatch, json.dumps({"sheet_name": "Receipt", "summary": "ok", "rows": [["Item", "Price"], ["Latte", 4.5]]}))

    result = runner.invoke(app, ["vision", str(img), "--purpose", "receipt", "--backend", "google", "--api-key", "k", "--output", str(out)])

    assert result.exit_code == 0, result.output
    payload = json.loads(result.stdout)
    assert payload["rows"] == [["Item", "Price"], ["Latte", 4.5]] and payload["saved"] == str(out)
    assert openpyxl.load_workbook(out)["Receipt"]["B2"].value == 4.5


# ─── calc ───────────────────────────────────────────────────────────────────


def test_calc_status_reports_unavailable_cleanly(monkeypatch):
    class Dummy:
        def __init__(self, *a, **kw):
            pass

        def is_available(self):
            return False

    monkeypatch.setattr("interfaces.cli._uno_driver", lambda settings: Dummy())

    result = runner.invoke(app, ["calc", "status"])

    assert result.exit_code == 0, result.output
    payload = json.loads(result.stdout)
    assert payload["available"] is False and "hint" in payload


def test_analyze_live_requires_available_calc(tmp_path: Path, monkeypatch):
    class Dummy:
        def is_available(self):
            return False

    monkeypatch.setattr("interfaces.cli._uno_driver", lambda settings: Dummy())

    result = runner.invoke(app, ["analyze", "--live", "q"])

    assert result.exit_code == 1
    assert "calc launch" in result.stderr


def test_analyze_live_applies_to_calc_with_yes(monkeypatch):
    from core.driver import ManifestSummary, WorkbookSnapshot

    class FakeUno:
        applied = []

        def is_available(self):
            return True

        def document_title(self):
            return "Budget.ods"

        def snapshot(self):
            return WorkbookSnapshot(sheets={"Sheet1": [["Name", "Score"], ["Ada", 91], ["Linus", 72]]}, source="Budget.ods")

        def execute_manifest(self, manifest):
            FakeUno.applied.append(manifest)
            return ManifestSummary(applied=1, errors=[])

    monkeypatch.setattr("interfaces.cli._uno_driver", lambda settings: FakeUno())

    result = runner.invoke(app, ["analyze", "--live", "Highlight rows that need review", "--yes"])

    assert result.exit_code == 0, result.output
    payload = json.loads(result.stdout)
    assert payload["applied"] == 1 and payload["target"] == "Budget.ods"
    assert len(FakeUno.applied) == 1


def test_analyze_live_port_option_reaches_uno_driver(monkeypatch):
    seen = {}

    class Dummy:
        def is_available(self):
            return False

    def fake_uno(settings):
        seen["port"] = settings.calc_port
        return Dummy()

    monkeypatch.setattr("interfaces.cli._uno_driver", fake_uno)

    result = runner.invoke(app, ["analyze", "--live", "--port", "2201", "anything"])

    assert result.exit_code == 1
    assert seen["port"] == 2201
    assert "port 2201" in result.stderr
